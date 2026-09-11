"""Validation: pyxform for the XlsForm itself, plus checks on the form folder."""

from __future__ import annotations

import shutil
from pathlib import Path

from .repo import Form


class ValidationError(Exception):
    """Raised when a built form is not a valid XlsForm."""


def pyxform_available() -> bool:
    try:
        import pyxform  # noqa: F401
    except ImportError:
        return False
    return True


def validate_xlsform(path: Path) -> list[str]:
    """Convert with pyxform. Returns warnings, raises ValidationError on failure.

    ODK Validate (a Java program bundled with pyxform) runs when `java` is on
    the PATH, which catches errors that the pure-Python conversion misses.
    """
    if not pyxform_available():
        raise ValidationError("pyxform is not installed. Run: pip install pyxform")

    from pyxform.errors import PyXFormError
    from pyxform.xls2xform import convert

    run_odk_validate = shutil.which("java") is not None
    try:
        result = convert(str(path), validate=run_odk_validate)
    except PyXFormError as e:
        raise ValidationError(str(e)) from None
    warnings = list(result.warnings or [])
    if not run_odk_validate:
        warnings.append("java not found, ODK Validate was skipped")
    return warnings


def check_form(form: Form) -> list[str]:
    """Static checks on a form folder that do not need pyxform. Returns problems."""
    from .build import QML_FILE_COLUMN, SURVEY_SHEET, check_media_references

    from openpyxl import load_workbook

    problems: list[str] = []
    if not form.xlsx.is_file():
        return [f"{form.xlsx} does not exist"]

    workbook = load_workbook(form.xlsx, read_only=True)
    if SURVEY_SHEET not in workbook.sheetnames:
        return [f"no '{SURVEY_SHEET}' sheet"]

    survey = workbook[SURVEY_SHEET]
    rows = survey.iter_rows(values_only=True)
    headers = list(next(rows, ()))
    if QML_FILE_COLUMN not in headers:
        problems.append(f"no '{QML_FILE_COLUMN}' column on the survey sheet")
    else:
        col = headers.index(QML_FILE_COLUMN)
        for row_number, row in enumerate(rows, start=2):
            name = row[col] if col < len(row) else None
            if name:
                try:
                    form.resolve_qml(str(name).strip())
                except FileNotFoundError as e:
                    problems.append(f"survey row {row_number}: {e}")

    if "settings" in workbook.sheetnames:
        settings_rows = list(workbook["settings"].iter_rows(values_only=True))
        if len(settings_rows) >= 2 and "namespaces" in settings_rows[0]:
            value = settings_rows[1][settings_rows[0].index("namespaces")]
            if not value or "ct=" not in str(value):
                problems.append("settings: 'namespaces' should declare ct=\"http://cybertracker.org/xforms\"")
        else:
            problems.append("settings: missing 'namespaces' column")
    else:
        problems.append("no 'settings' sheet")

    problems.extend(check_media_references(form))
    return problems
