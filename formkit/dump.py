"""Render a spreadsheet as Markdown tables so people and AI assistants can read it."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .build import QML_B64_COLUMN

MAX_CELL = 60


def dump_xlsx(path: Path, sheets: list[str] | None = None) -> str:
    workbook = load_workbook(path, read_only=True)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        if sheets and sheet.title not in sheets:
            continue
        chunks.append(f"## {sheet.title}\n")
        chunks.append(_table(sheet))
    return "\n".join(chunks)


def _table(sheet) -> str:
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    if not rows:
        return "_(empty)_\n"
    headers = rows[0]
    body = [r for r in rows[1:] if any(v is not None and str(v).strip() for v in r)]

    # Drop columns that are empty in every row so wide sheets stay readable.
    keep = [
        i
        for i, h in enumerate(headers)
        if h is not None and any(i < len(r) and r[i] is not None for r in body)
    ]
    if not keep:
        return "_(no data rows)_\n"

    def cell(value, header) -> str:
        if value is None:
            return ""
        text = str(value).replace("\n", " ").replace("|", "\\|")
        if header == QML_B64_COLUMN and len(text) > 16:
            return f"(embedded, {len(text)} chars)"
        if len(text) > MAX_CELL:
            return text[: MAX_CELL - 1] + "…"
        return text

    lines = [
        "| " + " | ".join(str(headers[i]) for i in keep) + " |",
        "| " + " | ".join("---" for _ in keep) + " |",
    ]
    for r in body:
        lines.append("| " + " | ".join(cell(r[i] if i < len(r) else None, headers[i]) for i in keep) + " |")
    return "\n".join(lines) + "\n"
