"""Turn a form folder into a finalized, self-contained XlsForm.

The source spreadsheet names its custom screens in a `bind::ct:content.qmlFile`
column. The build reads each referenced QML file, compresses it the way Qt
expects (zlib with a 4-byte big-endian length prefix), base64-encodes it, and
writes the result into a `bind::ct:content.qmlBase64z` column. The `qmlFile`
column is dropped, so the output needs nothing but its media files.
"""

from __future__ import annotations

import base64
import shutil
import struct
import zlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook

from .repo import Form

QML_FILE_COLUMN = "bind::ct:content.qmlFile"
QML_B64_COLUMN = "bind::ct:content.qmlBase64z"
SURVEY_SHEET = "survey"

Logger = Callable[[str], None]


class BuildError(Exception):
    """Raised when a form cannot be built."""


def qt_compress(data: bytes) -> bytes:
    """Compress the way QByteArray::qCompress does: 4-byte big-endian length + zlib."""
    return struct.pack(">I", len(data)) + zlib.compress(data)


def qt_decompress(data: bytes) -> bytes:
    """Inverse of qt_compress. Checks the length prefix."""
    if len(data) < 4:
        raise ValueError("qt-compressed data is too short")
    (expected_len,) = struct.unpack(">I", data[:4])
    out = zlib.decompress(data[4:])
    if len(out) != expected_len:
        raise ValueError(f"length prefix {expected_len} does not match payload {len(out)}")
    return out


def encode_qml(path: Path) -> str:
    """Read a QML file and return its qmlBase64z cell value."""
    text = path.read_text(encoding="utf-8")
    return base64.b64encode(qt_compress(text.encode("utf-8"))).decode("ascii")


def decode_qml(cell_value: str) -> str:
    """Turn a qmlBase64z cell value back into QML source."""
    return qt_decompress(base64.b64decode(cell_value)).decode("utf-8")


@dataclass
class BuildResult:
    form: Form
    output_dir: Path
    output_xlsx: Path
    qml: dict[str, Path] = field(default_factory=dict)  # file name -> resolved source
    media: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def build_form(
    form: Form,
    build_root: Path | None = None,
    validate: bool = True,
    log: Logger = print,
) -> BuildResult:
    """Build one form into build/<name>/ (or under *build_root*)."""
    if not form.xlsx.is_file():
        raise BuildError(f"{form.xlsx} does not exist")

    output_dir = form.build_dir(build_root)
    output_dir.mkdir(parents=True, exist_ok=True)
    result = BuildResult(form=form, output_dir=output_dir, output_xlsx=output_dir / form.xlsx.name)

    log(f"Building {form.id}")
    workbook = load_workbook(form.xlsx)
    if SURVEY_SHEET not in workbook.sheetnames:
        raise BuildError(f"{form.xlsx}: no '{SURVEY_SHEET}' sheet")

    survey = workbook[SURVEY_SHEET]
    headers = [cell.value for cell in survey[1]]
    if QML_FILE_COLUMN not in headers:
        raise BuildError(f"{form.xlsx}: no '{QML_FILE_COLUMN}' column on the survey sheet")
    qml_col = headers.index(QML_FILE_COLUMN)

    out = Workbook()
    out_survey = out.active
    out_survey.title = SURVEY_SHEET
    out_survey.append(_drop(headers, qml_col) + [QML_B64_COLUMN])

    for row_number, row in enumerate(survey.iter_rows(min_row=2, values_only=True), start=2):
        values = list(row) + [None] * (len(headers) - len(row))
        qml_name = values[qml_col]
        encoded = None
        if qml_name:
            qml_name = str(qml_name).strip()
            try:
                source = form.resolve_qml(qml_name)
            except FileNotFoundError as e:
                raise BuildError(f"{form.xlsx} row {row_number}: {e}") from None
            encoded = encode_qml(source)
            result.qml[qml_name] = source
            log(f"  embed {qml_name} <- {source.relative_to(form.root)}")
        out_survey.append(_drop(values, qml_col) + [encoded])

    for sheet_name in workbook.sheetnames:
        if sheet_name == SURVEY_SHEET:
            continue
        dst = out.create_sheet(title=sheet_name)
        for row in workbook[sheet_name].iter_rows(values_only=True):
            dst.append(list(row))

    out.save(result.output_xlsx)
    log(f"  wrote {result.output_xlsx.relative_to(form.root) if _inside(result.output_xlsx, form.root) else result.output_xlsx}")

    result.media = copy_media(form, output_dir)
    log(f"  copied {len(result.media)} media file(s)")

    result.warnings.extend(check_media_references(form))
    for warning in result.warnings:
        log(f"  warning: {warning}")

    if validate:
        from .validate import validate_xlsform

        pyxform_warnings = validate_xlsform(result.output_xlsx)
        for warning in pyxform_warnings:
            log(f"  pyxform: {warning}")
        result.warnings.extend(pyxform_warnings)
        log("  XlsForm is valid")

    return result


def copy_media(form: Form, output_dir: Path) -> list[str]:
    copied: list[str] = []
    for src in form.media_files():
        shutil.copyfile(src, output_dir / src.name)
        copied.append(src.name)
    return copied


def check_media_references(form: Form) -> list[str]:
    """Warn about `media::*` cells that name a file missing from media/."""
    available = {p.name for p in form.media_files()}
    problems: list[str] = []
    workbook = load_workbook(form.xlsx, read_only=True)
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            continue
        media_cols = [i for i, h in enumerate(headers) if isinstance(h, str) and h.startswith("media::")]
        if not media_cols:
            continue
        for row_number, row in enumerate(rows, start=2):
            for i in media_cols:
                value = row[i] if i < len(row) else None
                if value and str(value).strip() not in available:
                    problems.append(f"{sheet.title} row {row_number}: {headers[i]} '{value}' is not in {form.media_dir.name}/")
    return problems


def _drop(values: list, index: int) -> list:
    return values[:index] + values[index + 1 :]


def _inside(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False
