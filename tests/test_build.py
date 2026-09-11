"""Build every form in the repo and check the output is what CyberTracker needs."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from formkit.build import QML_B64_COLUMN, QML_FILE_COLUMN, build_form, decode_qml
from formkit.validate import check_form, pyxform_available


def _sheet_rows(path: Path, sheet: str) -> list[list]:
    ws = load_workbook(path, read_only=True)[sheet]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _records(rows: list[list]) -> list[dict]:
    """Rows as {header: value}, ignoring empty cells and unnamed columns."""
    headers = rows[0]
    out = []
    for row in rows[1:]:
        out.append({h: v for h, v in zip(headers, row) if h is not None and v is not None})
    return out


def test_form_passes_static_checks(form):
    assert check_form(form) == []


def test_build_embeds_qml_and_copies_media(form, tmp_path):
    result = build_form(form, build_root=tmp_path, validate=False, log=lambda _: None)

    assert result.output_xlsx.is_file()
    built = load_workbook(result.output_xlsx, read_only=True)
    source = load_workbook(form.xlsx, read_only=True)
    assert built.sheetnames == source.sheetnames

    src_rows = _sheet_rows(form.xlsx, "survey")
    out_rows = _sheet_rows(result.output_xlsx, "survey")
    assert QML_FILE_COLUMN not in out_rows[0]
    assert QML_B64_COLUMN in out_rows[0]

    # Every row keeps its other cells, and the embedded QML decodes to the file
    # that the resolution chain picks for it.
    assert len(src_rows) == len(out_rows)
    for src, out in zip(_records(src_rows), _records(out_rows)):
        qml_name = src.pop(QML_FILE_COLUMN, None)
        embedded = out.pop(QML_B64_COLUMN, None)
        assert src == out
        if qml_name:
            expected = form.resolve_qml(qml_name).read_text(encoding="utf-8")
            assert decode_qml(embedded) == expected
        else:
            assert embedded is None
    assert set(result.qml) == {r[QML_FILE_COLUMN] for r in _records(src_rows) if QML_FILE_COLUMN in r}
    assert len(result.qml) >= 1, "a form should embed at least one screen"

    # Other sheets are copied unchanged.
    for sheet in source.sheetnames:
        if sheet != "survey":
            assert _sheet_rows(form.xlsx, sheet) == _sheet_rows(result.output_xlsx, sheet)

    # Media is copied flat, byte for byte.
    expected_media = {p.name: p.read_bytes() for p in form.media_files()}
    assert set(result.media) == set(expected_media)
    for name, data in expected_media.items():
        assert (result.output_dir / name).read_bytes() == data
    assert result.warnings == [], result.warnings


@pytest.mark.skipif(not pyxform_available(), reason="pyxform not installed")
def test_built_form_is_valid_xlsform(form, tmp_path):
    result = build_form(form, build_root=tmp_path, validate=True, log=lambda _: None)
    assert result.output_xlsx.is_file()
